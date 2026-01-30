from __future__ import annotations

from pathlib import Path
from typing import List

from docx import Document

from .config import AppConfig
from .db import Book, Database
from .notifications import Notifier


class CompileWorkflow:
    def __init__(self, cfg: AppConfig, db: Database, notifier: Notifier) -> None:
        self.cfg = cfg
        self.db = db
        self.notifier = notifier

    def run(self) -> None:
        books: List[Book] = self.db.get_books_for_compilation()
        for book in books:
            try:
                self._process_book(book)
            except Exception as exc:  # noqa: BLE001
                self.notifier.notify_error("compile_stage", f"Book '{book.title}': {exc}")

    def _process_book(self, book: Book) -> None:
        # Gating logic:
        # Compile only if:
        #   final_review_notes_status = no_notes_needed OR
        #   final_review_notes exist
        status = (book.final_review_notes_status or "").strip().lower() if book.final_review_notes_status else ""
        has_notes = bool(book.final_review_notes and book.final_review_notes.strip())

        if not (status == "no_notes_needed" or has_notes):
            # Pause, not ready for compilation
            return

        chapters = self.db.get_chapters_for_book(book.id)
        if not chapters:
            # Nothing to compile
            return

        # Order by chapter_number already done in DB helper
        full_text_parts: List[str] = []
        for ch in chapters:
            title = ch.get("chapter_title") or f"Chapter {ch['chapter_number']}"
            content = ch.get("content") or ""
            full_text_parts.append(f"{title}\n\n{content}\n\n")

        full_text = "\n".join(full_text_parts).strip() + "\n"

        base_dir: Path = self.cfg.output.base_dir
        base_dir.mkdir(parents=True, exist_ok=True)

        txt_path = base_dir / f"{book.id}.txt"
        txt_path.write_text(full_text, encoding="utf-8")

        if self.cfg.output.generate_docx:
            self._write_docx(book_title=book.title, full_text=full_text, path=base_dir / f"{book.id}.docx")

        self.db.update_book_status(book.id, book_output_status="ready")
        self.notifier.notify_final_compiled(book.title, str(txt_path))

    @staticmethod
    def _write_docx(book_title: str, full_text: str, path: Path) -> None:
        doc = Document()
        doc.add_heading(book_title, level=1)
        for para in full_text.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para)
        doc.save(path)
